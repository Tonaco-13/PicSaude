#!/usr/bin/env python3
"""
importar_snapshot_tuss.py — gera `data/tuss_procedimentos.csv` e
`data/tuss_sigtap_mapeamento.csv` a partir da planilha OFICIAL da ANS.

ENG-027 §2. Roda OFFLINE, à mão, por um humano — nunca é chamado pela
aplicação, exatamente como o `importar_snapshot_sigtap.py`. O que a aplicação
lê é o CSV commitado; o que garante que o CSV não foi inventado é este script
mais o sha256 no MANIFEST.

    cd backend && python3 scripts/importar_snapshot_tuss.py \
        ../data/fontes-oficiais/tuss/padraotiss_mapeamento_tuss_sigtap.zip

O QUE A FONTE É — e por que ela resolve DOIS problemas de uma vez
------------------------------------------------------------------
`padraotiss_mapeamento_tuss_sigtap.zip` é publicado pela ANS na página
"Padrão TISS – Tabelas Relacionadas". Dentro dele, uma planilha com 5 abas,
das quais duas interessam:

  • **"TUSS 22 201606"** — a Tabela 22 (Terminologia de procedimentos e
    eventos em saúde) INTEIRA, com código, termo e vigências. É a fonte
    oficial que faltava: até aqui a casa tinha 35 procedimentos TUSS
    hardcoded em `app/ai/tuss_base.py`, sem fonte nenhuma.

  • **"Mapeamento ativos"** — o mapeamento TUSS→SIGTAP feito PELA ANS e PELO
    MS, com **grau de equivalência declarado** (1 a 5, legenda na aba
    "Tabela_Equivalência"). O ticket previa casar por nome normalizado; não é
    preciso, e não se deve: mapeamento oficial com grau declarado é melhor
    que heurística nossa, e não corre o risco do "biopsia ≈ biopsia" que a
    sabotagem da guarda persegue.

IDADE DA FONTE — DECLARADA, como manda a casa
----------------------------------------------
O arquivo é de **abril de 2017** e a Tabela 22 que ele carrega é a
competência **201606**. É a régua da ITU 2003 e da AMB 2008: a idade entra
declarada no campo `fonte` de TODA row, para que ninguém a leia como corrente.
A ANS mantém a TUSS viva (a consulta pública atual serve versões mais novas
por uma API que não respondeu a esta vantage — pendência registrada no
MANIFEST); quando uma extração mais nova pousar, este script roda de novo e o
`versao_snapshot` muda junto.
"""
from __future__ import annotations

import csv
import os
import sys
import zipfile

import openpyxl

_ABA_TUSS = "TUSS 22 201606"
_ABA_MAP = "Mapeamento ativos"
_VERSAO = "ANS-TISS 2017-04 (TUSS 22 competência 201606)"
_FONTE = ("ANS/TISS — Tabela 22 (Procedimentos e eventos em saúde), "
          "competência 201606, publicada no pacote de mapeamento de 2017-04")
_FONTE_MAP = ("ANS/TISS + MS — mapeamento oficial TUSS x SIGTAP, aba "
              "'Mapeamento ativos', pacote de 2017-04")

_GRAUS = {
    "1": "equivalência de significado, léxica e conceitual",
    "2": "equivalência de significado, com sinonímia",
    "3": "TUSS menos específico que o SIGTAP",
    "4": "TUSS mais específico que o SIGTAP",
    "5": "não é possível mapeamento",
    # A ANS escreve "dúvida" na coluna de grau em parte das linhas — é valor da
    # PRÓPRIA fonte, não sujeira de extração (a metodologia já avisa: "alguns
    # mapeamentos ainda não estão com o grau de equivalência atribuído"). Entra
    # como está: normalizar para vazio esconderia que a autoridade hesitou, e
    # essa hesitação é informação clínica para quem vai faturar.
    "dúvida": "a própria ANS registrou dúvida sobre a equivalência",
}


def _abrir(caminho_zip: str):
    with zipfile.ZipFile(caminho_zip) as z:
        nomes = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        if len(nomes) != 1:
            raise SystemExit(f"esperava 1 xlsx no zip, achei {nomes}")
        destino = os.path.join(os.path.dirname(caminho_zip) or ".", "_tuss_tmp.xlsx")
        with z.open(nomes[0]) as f, open(destino, "wb") as saida:
            saida.write(f.read())
    return destino, nomes[0]


def _data(v) -> str:
    """A planilha traz datetime; o CSV guarda ISO-8601 curto, ou vazio."""
    if not v:
        return ""
    return str(v)[:10]


def importar(caminho_zip: str, dir_saida: str) -> tuple[int, int]:
    xlsx, nome_interno = _abrir(caminho_zip)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    for aba in (_ABA_TUSS, _ABA_MAP):
        if aba not in wb.sheetnames:
            raise SystemExit(f"aba {aba!r} não existe em {nome_interno!r} — "
                             f"abas presentes: {wb.sheetnames}")

    # ── Tabela 22 ────────────────────────────────────────────────────────────
    # O cabeçalho não está na linha 1: a aba abre com linhas de título. Procuro
    # a linha que declara as colunas em vez de assumir posição — planilha
    # oficial muda de diagramação entre competências, e assumir linha fixa é
    # como assumir que a competência 202609 existe.
    ws = wb[_ABA_TUSS]
    linhas = ws.iter_rows(values_only=True)
    for row in linhas:
        if row and str(row[0] or "").strip() == "Código do Termo":
            break
    else:
        raise SystemExit(f"não achei o cabeçalho 'Código do Termo' na aba {_ABA_TUSS!r}")

    vistos: set[str] = set()
    procedimentos = []
    for row in linhas:
        codigo = str(row[0] or "").strip()
        termo = str(row[2] or "").strip()
        if not codigo or not termo or not codigo.isdigit():
            continue
        if codigo in vistos:      # a aba não deveria repetir; se repetir, a
            continue              # primeira ocorrência manda e o total acusa
        vistos.add(codigo)
        procedimentos.append({
            "codigo_tuss": codigo,
            "termo": termo,
            "vigencia_inicio": _data(row[3]),
            "vigencia_fim": _data(row[4]),
            "versao_snapshot": _VERSAO,
            "fonte": _FONTE,
        })

    # ── mapeamento oficial ───────────────────────────────────────────────────
    ws = wb[_ABA_MAP]
    linhas = ws.iter_rows(values_only=True)
    next(linhas)                              # cabeçalho, na primeira linha
    pares, vistos_par = [], set()
    for row in linhas:
        tuss = str(row[0] or "").strip()
        sigtap = str(row[2] or "").strip()
        if not tuss.isdigit() or not sigtap.isdigit():
            continue
        chave = (tuss, sigtap)
        if chave in vistos_par:
            continue
        vistos_par.add(chave)
        grau = str(row[5] or "").strip()
        pares.append({
            "codigo_tuss": tuss,
            "codigo_sigtap": sigtap,
            "termo_tuss": str(row[1] or "").strip(),
            "procedimento_sigtap": str(row[3] or "").strip(),
            "status": str(row[4] or "").strip(),
            "grau_equivalencia": grau,
            "grau_descricao": _GRAUS.get(grau, ""),
            "versao_snapshot": _VERSAO,
            "fonte": _FONTE_MAP,
        })

    os.makedirs(dir_saida, exist_ok=True)
    p1 = os.path.join(dir_saida, "tuss_procedimentos.csv")
    with open(p1, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(procedimentos[0].keys()),
                           lineterminator="\n")
        w.writeheader(); w.writerows(procedimentos)
    p2 = os.path.join(dir_saida, "tuss_sigtap_mapeamento.csv")
    with open(p2, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(pares[0].keys()), lineterminator="\n")
        w.writeheader(); w.writerows(pares)
    os.remove(xlsx)
    return len(procedimentos), len(pares)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    saida = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "..", "..", "data")
    n, m = importar(sys.argv[1], saida)
    print(f"tuss_procedimentos.csv: {n} procedimentos")
    print(f"tuss_sigtap_mapeamento.csv: {m} pares")
